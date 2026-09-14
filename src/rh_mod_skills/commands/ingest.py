"""rh-mod-skills ingest — register L1 sources with structure-preserving projection."""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

import click
from ruamel.yaml import YAML

from rh_mod_skills.common import (
    append_model_event,
    append_root_event,
    consumer_root,
    locked_update_tracking,
    log_info,
    log_warn,
    model_dir,
    now_iso,
    require_model,
    require_tracking,
    sha256_file,
    tracking_file,
)
from rh_mod_skills.pdf_tables import (
    SKIP_NONE_APPROVED,
    detect_pdf_tables,
    project_pdf_tables,
)

PLAN_NAME = "ingest-plan.yaml"


def _yaml() -> YAML:
    y = YAML()
    y.default_flow_style = False
    y.preserve_quotes = True
    return y


def ingest_plan_path(name: str) -> Path:
    return model_dir(name) / "process" / "plans" / PLAN_NAME


def source_name_from_path(path: Path) -> str:
    stem = path.stem.lower().replace("_", "-").replace(" ", "-")
    while "--" in stem:
        stem = stem.replace("--", "-")
    return stem.strip("-") or "source"


def detect_type(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return "excel"
    if suffix == ".csv":
        return "csv"
    if suffix == ".pdf":
        return "pdf"
    return "other"


def projection_intent(type_: str) -> str:
    return "tables" if type_ in {"excel", "csv"} else "skipped"


def load_plan(model: str) -> dict:
    path = ingest_plan_path(model)
    if not path.exists():
        raise click.ClickException(
            f"No ingest plan at {path}. Run `rh-mod-skills ingest plan {model}` first."
        )
    with open(path, encoding="utf-8") as f:
        return _yaml().load(f) or {}


def save_plan(model: str, data: dict) -> None:
    path = ingest_plan_path(model)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        _yaml().dump(data, f)


def _cell(value) -> str:
    if value is None:
        return ""
    return str(value)


def project_excel(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            sheets.append({"name": ws.title, "columns": [], "rows": []})
            continue
        columns = [_cell(c) for c in header]
        data_rows = []
        for row in rows_iter:
            values = [_cell(c) for c in row]
            if columns and len(values) < len(columns):
                values.extend([""] * (len(columns) - len(values)))
            data_rows.append(values[: len(columns)] if columns else values)
        sheets.append({"name": ws.title, "columns": columns, "rows": data_rows})
    wb.close()
    return sheets


def project_csv(path: Path) -> list[dict]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return [{"name": path.stem, "columns": [], "rows": []}]
    return [{"name": path.stem, "columns": rows[0], "rows": rows[1:]}]


def write_projection(dest: Path, original_name: str, type_: str, sheets: list[dict]) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    payload = {"source": original_name, "type": type_, "sheets": sheets}
    with open(dest, "w", encoding="utf-8") as f:
        _yaml().dump(payload, f)


def _consumer_relative(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(consumer_root()))
    except ValueError:
        return str(path)


def _register_source(tracking: dict, model_name: str, src_name: str, record: dict) -> None:
    m = require_model(tracking, model_name)
    sources = m.setdefault("sources", [])
    sources[:] = [s for s in sources if s.get("name") != src_name]
    sources.append(record)
    append_model_event(tracking, model_name, "source_added", f"Source registered: {src_name}")
    append_root_event(tracking, "source_added", f"{model_name}: {src_name}")


def _resolve_tracked_path(file_field: str) -> Path:
    path = Path(file_field)
    if path.is_absolute():
        return path
    return consumer_root() / file_field


def _has_structured_artifacts(md: Path) -> bool:
    structured = md / "structured"
    if not structured.is_dir():
        return False
    return any(p.is_file() for p in structured.rglob("*"))


@click.group()
def ingest():
    """Register L1 sources and write structure-preserving table projections."""


@ingest.command("plan")
@click.argument("model")
@click.option(
    "--source",
    "sources",
    multiple=True,
    required=True,
    type=click.Path(exists=True, dir_okay=False, path_type=Path),
    help="L1 file to register (repeatable)",
)
@click.option("--origin-url", default=None, help="Provenance URL for these sources")
def plan_cmd(model, sources, origin_url):
    """Write a draft ingest plan. Does not copy source files."""
    tracking = require_tracking()
    require_model(tracking, model)
    entries = []
    for src in sources:
        src = src.resolve()
        type_ = detect_type(src)
        entry = {
            "path": str(src),
            "name": source_name_from_path(src),
            "type": type_,
            "projection": projection_intent(type_),
            "origin_url": origin_url,
        }
        if type_ == "pdf":
            preview = detect_pdf_tables(src)
            entry["projection"] = preview.get("projection") or "skipped"
            if preview.get("skip_reason"):
                entry["skip_reason"] = preview["skip_reason"]
            if preview.get("tables") is not None:
                entry["tables"] = preview.get("tables") or []
            if preview.get("warnings"):
                entry["warnings"] = preview["warnings"]
        entries.append(entry)
    data = {"model": model, "status": "draft", "sources": entries}
    save_plan(model, data)
    log_info(f"Wrote ingest plan ({len(entries)} source(s), status=draft)")
    click.echo(f"  {ingest_plan_path(model)}")
    for entry in entries:
        click.echo(f"  - {entry['name']}  type={entry['type']}  projection={entry['projection']}")
        for table in entry.get("tables") or []:
            click.echo(
                f"      {table.get('id')}  {table.get('name')}  "
                f"pages={table.get('pages')}  rows={table.get('row_count')}  "
                f"decision={table.get('decision')}"
            )


@ingest.command("approve")
@click.argument("model")
def approve_cmd(model):
    """Mark the ingest plan approved (human gate)."""
    tracking = require_tracking()
    require_model(tracking, model)
    data = load_plan(model)
    data["status"] = "approved"
    save_plan(model, data)
    log_info(f"Approved ingest plan for {model}")


@ingest.command("implement")
@click.argument("model")
@click.option(
    "--acknowledge-drift",
    is_flag=True,
    help="Allow L1 replace when checksum changed; never deletes L2",
)
def implement_cmd(model, acknowledge_drift):
    """Copy sources and write table projections. Requires an approved plan."""
    tracking = require_tracking()
    require_model(tracking, model)
    plan = load_plan(model)
    if plan.get("status") != "approved":
        raise click.ClickException(
            f"Ingest plan is not approved (status={plan.get('status')!r}). "
            f"Run `rh-mod-skills ingest approve {model}` after review."
        )

    md = model_dir(model)
    raw_dir = md / "sources" / "raw"
    proj_dir = md / "sources" / "projections"
    raw_dir.mkdir(parents=True, exist_ok=True)
    proj_dir.mkdir(parents=True, exist_ok=True)

    existing = {s.get("name"): s for s in require_model(tracking, model).get("sources") or []}

    for entry in plan.get("sources") or []:
        src = Path(entry["path"])
        if not src.is_file():
            raise click.ClickException(f"Plan source missing on disk: {src}")
        name = entry["name"]
        type_ = entry["type"]
        dest = raw_dir / src.name
        new_sum = sha256_file(src)
        old = existing.get(name)
        checksum_changed = bool(old and old.get("checksum") and old.get("checksum") != new_sum)
        want_tables = entry.get("projection") == "tables"
        already_projected = bool(old and old.get("projection") == "tables" and dest.is_file())
        if old and old.get("checksum") == new_sum and dest.is_file() and not want_tables and not already_projected:
            log_info(f"Unchanged, skipping: {name}")
            continue
        if checksum_changed and _has_structured_artifacts(md) and not acknowledge_drift:
            raise click.ClickException(
                f"Source '{name}' checksum changed and structured/ has artifacts. "
                "Re-run with --acknowledge-drift to refresh L1 only (L2 is not overwritten)."
            )

        if not dest.is_file() or checksum_changed or old is None:
            shutil.copy2(src, dest)
        rec = {
            "name": name,
            "file": _consumer_relative(dest),
            "type": type_,
            "checksum": new_sum,
            "ingested_at": now_iso(),
            "origin_url": entry.get("origin_url"),
            "projection": entry.get("projection"),
        }

        included_pdf = [
            t for t in (entry.get("tables") or []) if t.get("decision") == "include"
        ]
        if type_ == "pdf" and want_tables and not included_pdf:
            want_tables = False
            rec["projection"] = "skipped"
            rec["skip_reason"] = SKIP_NONE_APPROVED

        if want_tables:
            if type_ == "excel":
                sheets = project_excel(dest)
            elif type_ == "csv":
                sheets = project_csv(dest)
            elif type_ == "pdf":
                sheets = project_pdf_tables(dest, included_pdf)
                if not sheets:
                    rec["projection"] = "skipped"
                    rec["skip_reason"] = SKIP_NONE_APPROVED
                    log_info(f"Registered {name} (projection skipped: {SKIP_NONE_APPROVED})")
                    locked_update_tracking(
                        lambda t, record=rec, model_name=model, src_name=name: _register_source(
                            t, model_name, src_name, record
                        )
                    )
                    continue
            else:
                raise click.ClickException(f"Cannot project type {type_} as tables")
            rec["projection"] = "tables"
            proj_path = proj_dir / f"{name}.yaml"
            write_projection(proj_path, src.name, type_, sheets)
            rec["projection_file"] = _consumer_relative(proj_path)
            n_rows = sum(len(s["rows"]) for s in sheets)
            log_info(f"Projected {name}: {len(sheets)} sheet(s), {n_rows} data row(s)")
        else:
            rec["projection"] = "skipped"
            rec["skip_reason"] = rec.get("skip_reason") or entry.get("skip_reason") or type_
            rec.pop("projection_file", None)
            stale = proj_dir / f"{name}.yaml"
            if stale.is_file():
                stale.unlink()
            log_info(f"Registered {name} (projection skipped: {rec['skip_reason']})")
            if type_ == "pdf":
                md_guess = dest.with_suffix(".md")
                if md_guess.exists():
                    log_warn("Markdown companion present; it is not the L1 projection")

        locked_update_tracking(
            lambda t, record=rec, model_name=model, src_name=name: _register_source(
                t, model_name, src_name, record
            )
        )


@ingest.command("verify")
@click.argument("model")
def verify_cmd(model):
    """Non-destructive ingest report. Does not write tracking events."""
    tracking = require_tracking()
    m = require_model(tracking, model)
    sources = m.get("sources") or []
    if not sources:
        raise click.ClickException(f"No sources registered for {model}. Run ingest implement first.")

    blocking = 0
    before = tracking_file().read_text(encoding="utf-8")
    for src in sources:
        name = src.get("name")
        file_field = src.get("file")
        path = _resolve_tracked_path(file_field)
        click.echo(f"{name}")
        if not path.is_file():
            click.echo("  original: missing")
            blocking += 1
            continue
        click.echo("  original: present")
        actual = sha256_file(path)
        expected = src.get("checksum")
        if actual != expected:
            click.echo("  checksum: DRIFT")
            blocking += 1
        else:
            click.echo("  checksum: match")
        proj = src.get("projection")
        if proj == "skipped":
            click.echo(f"  projection: skipped ({src.get('skip_reason', 'n/a')})")
        else:
            pfile = src.get("projection_file")
            ppath = _resolve_tracked_path(pfile) if pfile else None
            if not ppath or not ppath.is_file():
                click.echo("  projection: missing")
                blocking += 1
            else:
                with open(ppath, encoding="utf-8") as f:
                    data = _yaml().load(f)
                sheets = data.get("sheets") or []
                rows = sum(len(s.get("rows") or []) for s in sheets)
                cols = len(sheets[0].get("columns") or []) if sheets else 0
                click.echo(f"  projection: present  sheets={len(sheets)} columns={cols} rows={rows}")
    after = tracking_file().read_text(encoding="utf-8")
    if before != after:
        raise click.ClickException("verify mutated tracking.yaml — this is a bug")
    if blocking:
        raise click.ClickException(f"ingest verify failed: {blocking} blocking issue(s)")
    log_info("ingest verify passed")
